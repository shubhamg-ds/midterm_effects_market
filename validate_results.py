#!/usr/bin/env python3
"""Sanity checks for the supplied snapshot and several hand-built edge cases."""
from pathlib import Path
import argparse
import math
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from run_analysis import load_macro, event_study, wilson


def assert_close(x, y, tol=1e-9):
    assert abs(x-y) <= tol, f"{x} != {y}"


def fake_prices():
    cal=pd.bdate_range("2018-01-01",periods=280)
    bi=100
    close=np.full(len(cal),100.0)
    close[bi-18]=110.0
    close[bi-17]=90.0
    close[bi+1]=111.0      # new upside cross
    close[bi+2]=105.0      # re-enter [90,110]
    close[bi+3:]=112.0     # finish above breakout
    p=pd.DataFrame(index=cal)
    p["close"]=close;p["open"]=close;p["high"]=close+1;p["low"]=close-1
    p["valid_close"]=True;p["valid_intraday"]=True
    return cal,bi,p


def run_source_parity(data):
    market=pd.read_csv(data/"market_daily.csv",parse_dates=["date"],low_memory=False).set_index("date")
    for prefix, filename in [("dji","DJI_1D.csv"),("nasdaq","NASDAQ_DLY_IXIC_1D.csv"),("vix","CBOE_DLY_VIX_1D.csv")]:
        raw=pd.read_csv(data/filename,parse_dates=["time"]).set_index("time")["close"].sort_index()
        processed=market[prefix+"_close"].dropna().sort_index()
        assert raw.index.equals(processed.index), f"Date mismatch for {prefix}"
        assert np.allclose(raw.to_numpy(dtype=float),processed.to_numpy(dtype=float),rtol=0,atol=1e-9), f"Value mismatch for {prefix}"
    print("PASS raw-to-processed DJI, NASDAQ and VIX close parity")


def run_synthetic(data):
    macro=load_macro(data)
    cal,bi,p=fake_prices()
    ctx,h,paths,br,intra=event_study("SPX",2018,cal[bi],"Midterm",p,cal,macro,True)
    row=next(x for x in h if x["Horizon"]==5)
    assert row["First new cross"]=="Upside" and row["First cross session"]==1
    assert row["Endpoint available"] and row["Complete path"]
    detail=next(x for x in br if x["Break"]=="Upside" and x["Forward sessions"]==5)
    assert detail["Reentered range"] and detail["Finished above break"] and detail["Reentered then finished above"]
    assert_close(detail["Forward return"],112.0/111.0-1.0)
    # A price already above the high at the anchor does not create a new cross
    # until it re-enters and then recrosses.
    p2=p.copy();p2.loc[cal[bi],"close"]=111;p2.loc[cal[bi+1],"close"]=112
    p2.loc[cal[bi+2],"close"]=105;p2.loc[cal[bi+3],"close"]=111
    _,hh,_,_,_=event_study("SPX",2018,cal[bi],"Midterm",p2,cal,macro,True)
    row=next(x for x in hh if x["Horizon"]==5)
    assert row["Anchor state"]=="Above pre-high" and row["First cross session"]==3
    # One missing interior session invalidates the path, not a known exact
    # endpoint; it must not shift day 60 to another date.
    p3=p.copy();p3.loc[cal[bi+20],"valid_close"]=False
    _,hh,_,bb,_=event_study("SPX",2018,cal[bi],"Midterm",p3,cal,macro,True)
    row=next(x for x in hh if x["Horizon"]==60)
    assert row["Endpoint available"] and not row["Complete path"]
    assert_close(row["Return"],112/100-1)
    assert row["End date"]==cal[bi+60] and not bb
    assert_close(wilson(0,15)[1],0.2038833010358487,tol=1e-8)
    print("PASS synthetic crossing, recross, path-gap, exact-endpoint and Wilson tests")


def table_rows(ws):
    heads=[c.value for c in ws[4]]
    return [dict(zip(heads,vals)) for vals in ws.iter_rows(min_row=5,values_only=True) if any(v is not None for v in vals)]


def run_workbook(path):
    wb=load_workbook(path,data_only=True,read_only=True)
    post=table_rows(wb["Post horizons"])
    r=next(x for x in post if x["Asset"]=="SPX" and x["Sessions"]==120)
    assert r["Endpoint N"]==16 and r["Up count"]==16 and r["Path N"]==15
    assert r["Dip >=10% count"]==3 and r["Dip >=20% count"]==0
    assert_close(r["Up 95% low"],wilson(16,16)[0])
    b=table_rows(wb["After a break"])
    r=next(x for x in b if x["Asset"]=="SPX" and x["Break"]=="Upside" and x["Forward sessions"]==60)
    assert (r["N"],r["Reentered pre-range count"],r["Reentered AND finished above breakout count"])==(12,9,8)
    ledger=table_rows(wb["Event ledger"])
    r=next(x for x in ledger if x["Asset"]=="SPX" and x["Year"]==1994 and x["Horizon"]==60)
    assert r["Endpoint available"] and not r["Complete path"] and r["Return"] is not None
    assert "1994-12-30" in r["Exclusion detail"]
    q=next(x for x in ledger if x["Asset"]=="QQQ" and x["Year"]==2018 and x["Horizon"]==120)
    assert not q["Endpoint available"] and not q["Complete path"]
    for yr, expected in [(1974,-0.1345),(2002,-0.1253)]:
        rr=next(x for x in ledger if x["Asset"]=="SPX" and x["Year"]==yr and x["Horizon"]==120)
        assert abs(rr["Worst from anchor"]-expected)<0.0002 and rr["Return"]>0
    for d in table_rows(wb["Break directions"]):
        assert d["First up count"]+d["First down count"]+d["No new cross count"]==d["N"]
    rows=table_rows(wb["All probabilities"])
    for d in rows:
        k,n,p=d["Count"],d["N"],d["Frequency"]
        assert 0<=k<=n
        if n:
            assert_close(p,k/n)
            assert d["95% low"]-1e-12<=p<=d["95% high"]+1e-12
            actual_years=[] if not d["Success years"] else str(d["Success years"]).split(", ")
            assert len(actual_years)==k, (d["Asset"],d["Question"],k,actual_years)
        else:
            assert p is None and d["95% low"] is None and d["95% high"] is None
    crash=table_rows(wb["Crash risk"])
    pre=next(x for x in crash if x["Asset"]=="SPX" and x["Window"]=="At anchor")
    assert (pre["Count"],pre["N"],pre["Success years"])==(3,16,"1974, 2002, 2022")
    tail=next(x for x in crash if x["Asset"]=="SPX" and x["Sessions"]==120 and x["Question"]=="Fall >=20% from anchor at any close")
    assert tail["Count"]==0 and tail["N"]==15 and tail["95% high"]>.20
    print("PASS workbook sample counts, missing windows, cases, path partition and all probability rows")


if __name__=="__main__":
    parser=argparse.ArgumentParser()
    parser.add_argument("--workbook",type=Path,default=Path(__file__).resolve().parent/"Midterm_Market_Path_Study.xlsx")
    parser.add_argument("--data",type=Path,default=Path(__file__).resolve().parent/"data")
    args=parser.parse_args()
    run_source_parity(args.data)
    run_synthetic(args.data)
    run_workbook(args.workbook)

