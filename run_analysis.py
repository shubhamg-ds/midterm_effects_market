#!/usr/bin/env python3
"""Self-contained US midterm market-path study for the supplied project.

Place this file in the project root. It automatically reads market_daily.csv
and election_events.csv from data/processed/ and macro files from data/; the
older flat data/ layout also works. No network access or workbook.py is needed.
Run: python run_analysis.py
"""
from __future__ import annotations

import argparse
import math
from pathlib import Path

import numpy as np
import pandas as pd

HORIZONS = [5, 10, 20, 30, 60, 90, 120]
ASSETS = {"SPX": "spx", "DJI": "dji", "NASDAQ": "nasdaq", "SPY": "spy", "QQQ": "qqq"}
ETFS = {"SPY", "QQQ"}
MIN_ETF_MINUTES = 180  # conservative screen, not proof of an end-of-day close
BREAK_SEARCH = 60
TIGHT_RANGE = 0.05
PULLBACK = 0.03


def wilson(k: int, n: int, z: float = 1.959963984540054):
    """Two-sided 95% Wilson interval for a binomial proportion."""
    if not n:
        return (np.nan, np.nan)
    p = k / n
    den = 1 + z*z/n
    center = (p + z*z/(2*n))/den
    half = z * np.sqrt(p*(1-p)/n + z*z/(4*n*n)) / den
    return max(0.0, center-half), min(1.0, center+half)


def probability(values, years=None):
    """Return count, denominator, frequency, Wilson interval and contributing years.

    Missing values are removed, never silently turned into failures. `years` must
    be indexed in the same order as values. Empty samples stay missing.
    """
    v = pd.Series(values).reset_index(drop=True)
    mask = v.notna()
    vv = v[mask].astype(bool)
    n, k = int(len(vv)), int(vv.sum())
    lo, hi = wilson(k, n)
    ytext = ""
    if years is not None and n:
        yr = pd.Series(years).reset_index(drop=True)
        success_mask = v.map(lambda x: bool(x) if pd.notna(x) else False)
        ytext = ", ".join(str(int(y)) for y in yr[mask & success_mask])
    return dict(k=k, n=n, p=k/n if n else np.nan, lo=lo, hi=hi, years=ytext)


def nov_tuesday(year):
    """Tuesday following the first Monday in November, without assuming a session."""
    nov2 = pd.Timestamp(year=year, month=11, day=2)
    return nov2 + pd.Timedelta(days=(1 - nov2.weekday()) % 7)


def locate_input(data_dir: Path, filename: str, *, processed: bool = False) -> Path:
    """Locate a required input in the user's project or the original flat archive.

    The preferred Windows project layout is data/processed/{market_daily,
    election_events}.csv plus the macro CSV files directly under data/. The
    script also supports the previous flat data/ layout, and accepts --data
    pointing to the project root, the data folder, or the processed folder.
    All paths are pathlib paths and work regardless of the current directory.
    """
    base = Path(data_dir).expanduser().resolve()
    if processed:
        candidates = [base / "processed" / filename, base / filename,
                      base / "data" / "processed" / filename, base / "data" / filename]
    else:
        candidates = [base / filename, base / "data" / filename,
                      base / "processed" / filename, base / "data" / "processed" / filename]
        if base.name.lower() == "processed":
            candidates.insert(0, base.parent / filename)
    unique = list(dict.fromkeys(candidates))
    for path in unique:
        if path.is_file():
            print(f"  {filename}: {path}")
            return path
    tried = "\n  ".join(str(path) for path in unique)
    raise FileNotFoundError(
        f"Required file '{filename}' was not found. Looked in:\n  {tried}\n"
        "Use --data to point to the project's data folder, or check the file name."
    )


def load_data(data_dir: Path):
    market_file = locate_input(data_dir, "market_daily.csv", processed=True)
    elections_file = locate_input(data_dir, "election_events.csv", processed=True)
    market = pd.read_csv(market_file, parse_dates=["date"], low_memory=False)
    events = pd.read_csv(elections_file, parse_dates=["election_date", "first_market_session_after"])
    assert not market["date"].isna().any(), "Bad market date"
    assert not market["date"].duplicated().any(), "Duplicate dates in market_daily.csv"
    assert not events["election_date"].duplicated().any(), "Duplicate event dates"
    events = events[events.event_type.astype(str).str.contains("midterm", case=False, na=False)].copy()
    events = events.sort_values("election_date").reset_index(drop=True)
    if events.empty:
        raise ValueError("No rows labeled midterm in election_events.csv")
    assert all(nov_tuesday(int(r.election_year)) == r.election_date for r in events.itertuples()), "Unexpected midterm date"
    market = market.sort_values("date").set_index("date")
    for key in [f"{p}_{c}" for p in ASSETS.values() for c in ["open", "high", "low", "close"]]:
        market[key] = pd.to_numeric(market[key], errors="coerce")
    for p in ["spy", "qqq"]:
        market[f"{p}_n_minutes"] = pd.to_numeric(market[f"{p}_n_minutes"], errors="coerce")
    # The DJI daily series supplies the reference session calendar. SPX alone
    # contains extra holiday/closure observations, e.g. Thanksgiving 1982 and
    # 2001-09-11; a union calendar would create false post-election sessions.
    cal = market.index[market["dji_close"].notna() & (market["dji_close"] > 0)]
    assert cal.is_monotonic_increasing and not cal.has_duplicates
    cal_mismatch = []
    for ev in events.itertuples():
        j = cal.searchsorted(ev.election_date, side="right")
        nextday = cal[j] if j < len(cal) else pd.NaT
        if pd.notna(ev.first_market_session_after) and nextday != ev.first_market_session_after:
            cal_mismatch.append((ev.election_year, str(nextday), str(ev.first_market_session_after)))
    if cal_mismatch:
        raise ValueError(f"DJI reference calendar disagrees with election file: {cal_mismatch}")
    px = {}
    for asset, prefix in ASSETS.items():
        cols = [f"{prefix}_{x}" for x in ["open", "high", "low", "close"]]
        if asset in ETFS:
            cols.append(f"{prefix}_n_minutes")
        p = market.reindex(cal)[cols].copy()
        p.columns = ["open", "high", "low", "close"] + (["minutes"] if asset in ETFS else [])
        p["valid_close"] = np.isfinite(p.close) & (p.close > 0)
        if asset in ETFS:
            # 180 does not guarantee the final minute is present. It does screen
            # obvious truncations while allowing normal ~210-minute early closes.
            p["valid_close"] &= p.minutes.notna() & (p.minutes >= MIN_ETF_MINUTES)
        ohlc_finite = p[["open", "high", "low", "close"]].notna().all(axis=1)
        p["valid_intraday"] = (p.valid_close & ohlc_finite & (p.high > p.low)
                               & (p.high >= p.close) & (p.low <= p.close)
                               & (p.high >= p.open) & (p.low <= p.open))
        px[asset] = p
    return market, events, cal, px


def load_macro(data_dir):
    macro = {}
    for name, fname, datecol, valcol in [
        ("vix", "CBOE_DLY_VIX_1D.csv", "time", "close"),
        ("dff", "DFF.csv", "observation_date", "DFF"),
        ("dgs2", "DGS2.csv", "observation_date", "DGS2"),
        ("dgs10", "DGS10.csv", "observation_date", "DGS10"),
        ("cpi", "CPIAUCSL.csv", "observation_date", "CPIAUCSL"),
    ]:
        macro_file = locate_input(data_dir, fname)
        df = pd.read_csv(macro_file, parse_dates=[datecol])
        df[valcol] = pd.to_numeric(df[valcol], errors="coerce")
        df = df.dropna(subset=[datecol, valcol]).sort_values(datecol)
        df = df.drop_duplicates(datecol, keep="last").set_index(datecol)
        macro[name] = df[valcol]
    # Observation date + 60 calendar days is an approximation, not a historical
    # release/vintage database. CPI revisions remain embedded in this file.
    cpi = macro["cpi"]
    yoy = cpi / cpi.shift(12) - 1
    yoy.index = yoy.index + pd.Timedelta(days=60)
    macro["cpi_yoy_60d"] = yoy.dropna()
    return macro


def last_at(series, date, max_age_days=None):
    if series.empty:
        return np.nan, pd.NaT
    i = series.index.searchsorted(date, side="right") - 1
    if i < 0:
        return np.nan, pd.NaT
    d = series.index[i]
    if max_age_days is not None and (date-d).days > max_age_days:
        return np.nan, pd.NaT
    return float(series.iloc[i]), d


def macro_at(macro, anchor):
    if pd.isna(anchor):
        return {}
    vix, vixdate = last_at(macro["vix"], anchor, 7)
    # Rates are taken strictly before the anchor close; publication and revision
    # histories are unknown and still prevent a point-in-time trading backtest.
    prior_day = anchor - pd.Timedelta(days=1)
    dff, dffdate = last_at(macro["dff"], prior_day, 7)
    d2, d2date = last_at(macro["dgs2"], prior_day, 7)
    d10, d10date = last_at(macro["dgs10"], prior_day, 7)
    curve = (d10-d2) if pd.notna(d2date) and d2date == d10date else np.nan
    cpi, cpi_eff = last_at(macro["cpi_yoy_60d"], anchor, 125)
    return {"VIX close": vix, "VIX date": vixdate, "Fed funds %": dff,
            "Fed funds obs": dffdate, "2Y yield %": d2, "10Y yield %": d10,
            "Yield obs": d2date if pd.notna(curve) else pd.NaT,
            "10Y-2Y (pp)": curve, "CPI YoY": cpi,
            "CPI effective date": cpi_eff,
            "CPI observation month": cpi_eff - pd.Timedelta(days=60) if pd.notna(cpi_eff) else pd.NaT}


def event_study(asset, year, edate, cohort, p, cal, macro, detailed=True):
    """Build exact-session event, path, breakout and intraday records.

    Returns are close-to-close price changes. Window validity is checked against
    the shared DJI calendar BEFORE dropping anything; absent rows are not compressed.
    """
    bi = cal.searchsorted(edate, side="right") - 1
    if bi < 0:
        return None, [], [], [], []
    anchor = cal[bi]
    basevalid = bool(p.valid_close.iloc[bi])
    base = float(p.close.iloc[bi]) if basevalid else np.nan
    before = p.iloc[max(0, bi-20):bi]
    prevalid = len(before) == 20 and bool(before.valid_close.all())
    hi = float(before.close.max()) if prevalid else np.nan
    lo = float(before.close.min()) if prevalid else np.nan
    pre_ret = base / float(before.close.iloc[0]) - 1 if prevalid and basevalid else np.nan
    anchor_state = ("Above pre-high" if base > hi else "Below pre-low" if base < lo else "Inside range") if prevalid and basevalid else "Unavailable"
    trailing = p.iloc[max(0, bi-251):bi+1]
    trailvalid = len(trailing) == 252 and bool(trailing.valid_close.all())
    trailpeak = float(trailing.close.max()) if trailvalid else np.nan
    peakgap = base / trailpeak - 1 if trailvalid and basevalid else np.nan
    # 20 pre-event close-to-close returns, including the move into the anchor.
    premove = pd.concat([before.close, p.close.iloc[bi:bi+1]]) if prevalid and basevalid else pd.Series(dtype=float)
    vol = float(np.log(premove / premove.shift(1)).dropna().std(ddof=1) * np.sqrt(252)) if len(premove) == 21 else np.nan
    ctx = {"Cohort": cohort, "Asset": asset, "Year": int(year), "Election date": edate,
           "Anchor date": anchor, "Anchor close": base, "Anchor state": anchor_state,
           "Pre-20 close high": hi, "Pre-20 close low": lo, "Pre-20 return": pre_ret,
           "252-session peak": trailpeak, "Gap to 252-session peak": peakgap,
           "Pre-20 realized vol": vol, "Pre-20 complete": bool(prevalid),
           "252-session complete": bool(trailvalid), **macro_at(macro, anchor)}
    horizon_rows, path_rows, break_rows, intraday_rows = [], [], [], []
    
    # Pre-event intraday levels exclude the anchor, exactly like the notebook.
    valid_intraday_pre = len(before) == 20 and bool(before.valid_intraday.all())
    intra_hi = float(before.high.max()) if valid_intraday_pre else np.nan
    intra_lo = float(before.low.min()) if valid_intraday_pre else np.nan
    for h in HORIZONS:
        future = p.iloc[bi+1:bi+1+h]
        complete = bool(basevalid and len(future) == h and future.valid_close.all())
        # The endpoint is knowable even when an INTERMEDIATE close is missing.
        # Only path-dependent quantities require every intervening session.
        endpoint_ok = bool(basevalid and len(future) == h and future.valid_close.iloc[-1])
        missing = []
        if not basevalid:
            missing.append(f"anchor {anchor.date()} invalid/missing")
        if len(future) < h:
            missing.append(f"calendar ends after {len(future)} of {h} sessions")
        for d, row in future.loc[~future.valid_close].iterrows():
            why = "insufficient minute bars" if (asset in ETFS and pd.notna(row.get("close")) and pd.notna(row.get("minutes")) and row.get("minutes") < MIN_ETF_MINUTES) else "missing/invalid close"
            missing.append(f"{d.date()} {why}")
        rec = {"Cohort": cohort, "Asset": asset, "Year": int(year), "Election date": edate,
               "Anchor date": anchor, "Horizon": h, "End date": cal[bi+h] if bi+h < len(cal) else pd.NaT,
               "Endpoint available": endpoint_ok, "Complete path": complete, "Exclusion detail": "; ".join(missing),
               "Anchor close": base, "Anchor state": anchor_state,
               "Pre-20 close high": hi, "Pre-20 close low": lo,
               "Return": np.nan, "Worst from anchor": np.nan, "Max drawdown": np.nan,
               "Close range": np.nan, "Tight range": np.nan, "Above pre-high": np.nan,
               "Below pre-low": np.nan, "New upside cross": np.nan, "New downside cross": np.nan,
               "First new cross": np.nan, "First cross session": np.nan}
        if endpoint_ok:
            rec["Return"] = float(future.close.iloc[-1]/base - 1)
        if complete:
            arr = np.r_[base, future.close.to_numpy(dtype=float)]
            ret = arr[-1]/base - 1
            worst = float(arr.min()/base - 1)
            dd = float(np.min(arr/np.maximum.accumulate(arr)-1))
            spread = float(arr.max()/arr.min()-1)
            rec.update({"Return": ret, "Worst from anchor": worst, "Max drawdown": dd,
                        "Close range": spread, "Tight range": spread <= TIGHT_RANGE})
            if prevalid:
                ups = (arr[:-1] <= hi) & (arr[1:] > hi)
                downs = (arr[:-1] >= lo) & (arr[1:] < lo)
                hits = np.flatnonzero(ups | downs)
                first = int(hits[0]) if len(hits) else None
                rec.update({"Above pre-high": bool(np.any(arr[1:] > hi)),
                            "Below pre-low": bool(np.any(arr[1:] < lo)),
                            "New upside cross": bool(np.any(ups)), "New downside cross": bool(np.any(downs)),
                            "First new cross": ("Upside" if ups[first] else "Downside") if first is not None else "None",
                            "First cross session": first+1 if first is not None else np.nan})
        horizon_rows.append(rec)
        if detailed and cohort == "Midterm":
            # All pre and post bars must have genuine high/low dispersion. In
            # particular NASDAQ 1974-82 is close-only disguised as OHLC.
            intra_ok = bool(complete and valid_intraday_pre and len(future) == h and future.valid_intraday.all())
            irow = {"Asset": asset, "Year": int(year), "Horizon": h, "Eligible": intra_ok,
                    "Prior intraday high": intra_hi, "Prior intraday low": intra_lo,
                    "High touched": np.nan, "Low touched": np.nan,
                    "High close-confirmed": np.nan, "Low close-confirmed": np.nan,
                    "First intraday touch": np.nan}
            if intra_ok:
                ups = future.high.to_numpy() > intra_hi
                dns = future.low.to_numpy() < intra_lo
                firsts = np.flatnonzero(ups | dns)
                if len(firsts):
                    k = firsts[0]
                    direction = "Both same day—order unknown" if ups[k] and dns[k] else ("High" if ups[k] else "Low")
                else:
                    direction = "None"
                irow.update({"High touched": bool(ups.any()), "Low touched": bool(dns.any()),
                             "High close-confirmed": bool((future.close > intra_hi).any()),
                             "Low close-confirmed": bool((future.close < intra_lo).any()),
                             "First intraday touch": direction})
            intraday_rows.append(irow)
    if detailed and cohort == "Midterm" and basevalid:
        for i in range(max(0, bi-20), min(len(cal), bi+121)):
            r = p.iloc[i]
            path_rows.append({"Asset": asset, "Year": int(year), "Session": i-bi, "Date": cal[i],
                              "Open": r.open, "High": r.high, "Low": r.low, "Close": r.close,
                              "Minute bars": r.get("minutes", np.nan), "Usable close": bool(r.valid_close),
                              "Close / anchor": r.close/base if r.valid_close else np.nan,
                              "Pre-high / anchor": hi/base if prevalid else np.nan,
                              "Pre-low / anchor": lo/base if prevalid else np.nan})
    # Only events with a completely observed first 60-session search period
    # qualify. First upside and first downside crossings are analyzed separately.
    search = p.iloc[bi+1:bi+1+BREAK_SEARCH]
    if detailed and cohort == "Midterm" and basevalid and prevalid and len(search) == BREAK_SEARCH and search.valid_close.all():
        seq = np.r_[base, search.close.to_numpy(dtype=float)]
        ups = (seq[:-1] <= hi) & (seq[1:] > hi)
        dns = (seq[:-1] >= lo) & (seq[1:] < lo)
        for side, hits in [("Upside", ups), ("Downside", dns)]:
            ii = np.flatnonzero(hits)
            if not len(ii):
                continue
            session = int(ii[0])+1
            br_i = bi+session
            br_close = float(p.close.iloc[br_i])
            for h in HORIZONS:
                after = p.iloc[br_i+1:br_i+1+h]
                ok = bool(len(after)==h and after.valid_close.all())
                rr = {"Asset": asset, "Year": int(year), "Break": side,
                      "Election date": edate, "Anchor state": anchor_state,
                      "Break session": session, "Break date": cal[br_i], "Break close": br_close,
                      "Pre-20 close high": hi, "Pre-20 close low": lo, "Forward sessions": h,
                      "Forward end date": cal[br_i+h] if br_i+h < len(cal) else pd.NaT,
                      "Forward complete": ok, "Forward return": np.nan,
                      "Worst after break": np.nan, "Reentered range": np.nan,
                      "Pullback >=3%": np.nan, "Finished above break": np.nan,
                      "Reentered then finished above": np.nan,
                      "Pullback then finished above": np.nan, "Directional continuation >=3%": np.nan,
                      "Tight range": np.nan}
                if ok:
                    arr = np.r_[br_close, after.close.to_numpy(dtype=float)]
                    ret = float(arr[-1]/br_close-1)
                    worst = float(arr.min()/br_close-1)
                    reenter = bool(((after.close >= lo) & (after.close <= hi)).any())
                    fell = worst <= -PULLBACK
                    up = ret > 0
                    continuation = ret >= .03 if side == "Upside" else ret <= -.03
                    rr.update({"Forward return": ret, "Worst after break": worst,
                               "Reentered range": reenter, "Pullback >=3%": bool(fell),
                               "Finished above break": bool(up), "Reentered then finished above": bool(reenter and up),
                               "Pullback then finished above": bool(fell and up),
                               "Directional continuation >=3%": bool(continuation),
                               "Tight range": bool(arr.max()/arr.min()-1 <= TIGHT_RANGE)})
                break_rows.append(rr)
    return ctx, horizon_rows, path_rows, break_rows, intraday_rows


def pre_election_study(asset, year, edate, p, cal):
    """Pre-anchor close-to-close returns on exactly h common trading sessions."""
    bi = cal.searchsorted(edate, side="right") - 1
    rows=[]
    for h in HORIZONS:
        start = bi-h
        w = p.iloc[start:bi+1] if start >= 0 else p.iloc[0:0]
        ok = len(w)==h+1 and bool(w.valid_close.all())
        endpoint_ok = len(w)==h+1 and bool(w.valid_close.iloc[0]) and bool(w.valid_close.iloc[-1])
        r={"Asset":asset, "Year":int(year), "Election date":edate, "Sessions":h,
           "Start date":cal[start] if start>=0 else pd.NaT,
           "Anchor date":cal[bi] if bi>=0 else pd.NaT, "Endpoints available":endpoint_ok, "Complete path":ok,
           "Return":np.nan, "Max drawdown":np.nan}
        if endpoint_ok:
            r["Return"]=float(w.close.iloc[-1]/w.close.iloc[0]-1)
        if ok:
            arr=w.close.to_numpy(dtype=float)
            r["Max drawdown"]=float(np.min(arr/np.maximum.accumulate(arr)-1))
        rows.append(r)
    return rows


def build_summaries(ctx, evh, br, intr, compare, events, pre):
    """Make concise, auditable tables; keep counts and years next to frequencies."""
    summary, patterns, breaksum, intra_sum, allprob, regime, benchmark, pre_summary = [], [], [], [], [], [], [], []
    metrics = [
        ("Finish above anchor", lambda x: x["Return"].map(lambda v: v > 0 if pd.notna(v) else np.nan)),
        ("Finish at least +3%", lambda x: x["Return"].map(lambda v: v >= .03 if pd.notna(v) else np.nan)),
        ("Finish within +/-2%", lambda x: x["Return"].map(lambda v: abs(v) <= .02 if pd.notna(v) else np.nan)),
        ("Finish at least -3%", lambda x: x["Return"].map(lambda v: v <= -.03 if pd.notna(v) else np.nan)),
        ("Fall >=5% from anchor at any close", lambda x: x["Worst from anchor"].map(lambda v: v <= -.05 if pd.notna(v) else np.nan)),
        ("Fall >=10% from anchor at any close", lambda x: x["Worst from anchor"].map(lambda v: v <= -.10 if pd.notna(v) else np.nan)),
        ("Fall >=20% from anchor at any close", lambda x: x["Worst from anchor"].map(lambda v: v <= -.20 if pd.notna(v) else np.nan)),
        ("Peak-to-trough drawdown >=10%", lambda x: x["Max drawdown"].map(lambda v: v <= -.10 if pd.notna(v) else np.nan)),
        ("Tight: whole close range <=5%", lambda x: x["Tight range"]),
        ("Any close above pre-20 closing high", lambda x: x["Above pre-high"]),
        ("Any close below pre-20 closing low", lambda x: x["Below pre-low"]),
        ("First NEW crossing: upside", lambda x: x["First new cross"].map(lambda v: v == "Upside" if pd.notna(v) else np.nan)),
        ("First NEW crossing: downside", lambda x: x["First new cross"].map(lambda v: v == "Downside" if pd.notna(v) else np.nan)),
        ("No NEW range crossing", lambda x: x["First new cross"].map(lambda v: v == "None" if pd.notna(v) else np.nan)),
        ("NEW upside crossing at least once", lambda x: x["New upside cross"]),
        ("NEW downside crossing at least once", lambda x: x["New downside cross"]),
    ]
    for asset in ASSETS:
        for h in HORIZONS:
            preg = pre[(pre.Asset==asset) & (pre.Sessions==h)].copy().reset_index(drop=True)
            pres = probability(preg["Return"].map(lambda v: v > 0 if pd.notna(v) else np.nan), preg.Year)
            pre_summary.append({"Asset":asset,"Sessions before":h,"Endpoint N":pres["n"], "Up count":pres["k"],
                                "Up frequency":pres["p"],"95% low":pres["lo"],"95% high":pres["hi"],
                                "Median return":preg["Return"].median(),"Path N":int(preg["Max drawdown"].notna().sum()),
                                "Median max drawdown":preg["Max drawdown"].median()})
            allprob.append({"Section":"Before election", "Asset":asset, "Condition":"All valid midterms", "Horizon":h,
                            "Question":"Finish above the close h sessions before anchor", "Count":pres["k"],"N":pres["n"],
                            "Frequency":pres["p"],"95% low":pres["lo"],"95% high":pres["hi"],"Success years":pres["years"]})
            g = evh[(evh.Asset==asset) & (evh.Horizon==h)].copy().reset_index(drop=True)
            pp = {}
            for label, fn in metrics:
                res = probability(fn(g), g.Year)
                pp[label] = res
                allprob.append({"Section":"Election horizon", "Asset":asset, "Condition":"All valid midterms", "Horizon":h,
                                "Question":label, "Count":res["k"], "N":res["n"], "Frequency":res["p"],
                                "95% low":res["lo"], "95% high":res["hi"], "Success years":res["years"]})
            r = g["Return"].dropna()
            worst = g["Worst from anchor"].dropna()
            positive = pp["Finish above anchor"]
            dip5 = pp["Fall >=5% from anchor at any close"]
            dip10 = pp["Fall >=10% from anchor at any close"]
            dip20 = pp["Fall >=20% from anchor at any close"]
            tight = pp["Tight: whole close range <=5%"]
            summary.append({"Asset":asset, "Sessions":h, "Endpoint N":positive["n"], "Up count":positive["k"], "Up frequency":positive["p"],
                            "Up 95% low":positive["lo"], "Up 95% high":positive["hi"],
                            "Median return":r.median() if len(r) else np.nan,
                            "Lowest return":r.min() if len(r) else np.nan,
                            "Path N":dip5["n"], "Median worst dip":worst.median() if len(worst) else np.nan,
                            "Dip >=5% count":dip5["k"], "Dip >=5% frequency":dip5["p"],
                            "Dip >=10% count":dip10["k"], "Dip >=10% frequency":dip10["p"],
                            "Dip >=20% count":dip20["k"], "Dip >=20% frequency":dip20["p"],
                            "Tight count":tight["k"], "Tight frequency":tight["p"]})
            a,b,c = (pp["First NEW crossing: upside"], pp["First NEW crossing: downside"], pp["No NEW range crossing"])
            ab,bel = pp["Any close above pre-20 closing high"],pp["Any close below pre-20 closing low"]
            patterns.append({"Asset":asset, "Sessions":h, "N":a["n"], "First up count":a["k"], "First up frequency":a["p"],
                             "First down count":b["k"], "First down frequency":b["p"], "No new cross count":c["k"],
                             "No new cross frequency":c["p"], "Any above count":ab["k"], "Any above frequency":ab["p"],
                             "Any below count":bel["k"], "Any below frequency":bel["p"]})
            # Same-calendar odd-year benchmark, descriptive and NOT a causal control.
            cg = compare[(compare.Asset==asset) & (compare.Horizon==h)].copy().reset_index(drop=True)
            for q, fn in [metrics[0], metrics[4], metrics[5], metrics[8], metrics[11]]:
                mm, cc = probability(fn(g),g.Year), probability(fn(cg),cg.Year)
                benchmark.append({"Asset":asset, "Sessions":h, "Question":q,
                                  "Midterm count":mm["k"], "Midterm N":mm["n"], "Midterm frequency":mm["p"],
                                  "Midterm 95% low":mm["lo"], "Midterm 95% high":mm["hi"],
                                  "Odd-year count":cc["k"], "Odd-year N":cc["n"], "Odd-year frequency":cc["p"],
                                  "Odd-year 95% low":cc["lo"], "Odd-year 95% high":cc["hi"],
                                  "Difference (pp)":(mm["p"]-cc["p"])*100 if mm["n"] and cc["n"] else np.nan})
            ig = intr[(intr.Asset==asset) & (intr.Horizon==h)].copy().reset_index(drop=True)
            imetrics = [("Intraday high touched", "High touched"), ("Intraday low touched", "Low touched"),
                        ("Close above pre-intraday high", "High close-confirmed"),
                        ("Close below pre-intraday low", "Low close-confirmed")]
            irow = {"Asset":asset, "Sessions":h}
            for label,col in imetrics:
                res = probability(ig[col], ig.Year)
                irow["N"] = res["n"]
                irow[label+" count"] = res["k"]
                irow[label+" frequency"] = res["p"]
                allprob.append({"Section":"Intraday touch", "Asset":asset, "Condition":"Genuine pre/post OHLC required", "Horizon":h,
                                "Question":label, "Count":res["k"], "N":res["n"], "Frequency":res["p"],
                                "95% low":res["lo"], "95% high":res["hi"], "Success years":res["years"]})
            intra_sum.append(irow)
    # Forward outcomes are conditioned on a breakout and on the additional complete
    # forward observation period. Horizons start at the BREAKOUT close, not election.
    bmetrics = [("Reentered pre-range", "Reentered range"), ("Fell >=3% from breakout", "Pullback >=3%"),
                ("Finished above breakout", "Finished above break"),
                ("Reentered AND finished above breakout", "Reentered then finished above"),
                ("Fell >=3% AND finished above breakout", "Pullback then finished above"),
                ("Directional continuation >=3%", "Directional continuation >=3%"),
                ("Tight close range <=5%", "Tight range")]
    for asset in ASSETS:
        for side in ["Upside", "Downside"]:
            for h in HORIZONS:
                g = br[(br.Asset==asset) & (br.Break==side) & (br["Forward sessions"]==h)].copy().reset_index(drop=True)
                row = {"Asset":asset, "Break":side, "Forward sessions":h}
                for label,col in bmetrics:
                    res = probability(g[col],g.Year)
                    row["N"] = res["n"]
                    row[label+" count"] = res["k"]
                    row[label+" frequency"] = res["p"]
                    allprob.append({"Section":"After first NEW break", "Asset":asset,
                                    "Condition":side+" close crossing within first 60 sessions", "Horizon":h,
                                    "Question":label, "Count":res["k"], "N":res["n"], "Frequency":res["p"],
                                    "95% low":res["lo"], "95% high":res["hi"], "Success years":res["years"]})
                breaksum.append(row)
    # Exploratory SPX conditions; all thresholds are explicit, but some are
    # judgment calls and the 16 historical events make split samples very small.
    c = ctx[(ctx.Asset=="SPX") & (ctx.Cohort=="Midterm")].copy()
    conditions = [
        ("Pre-20 return", "Pre-20 return", lambda v: v >= 0, "Non-negative", "Negative"),
        ("Underwater vs 252-session peak", "Gap to 252-session peak", lambda v: v <= -.20, "At least 20% below peak", "Less than 20% below peak"),
        ("Pre-20 realized volatility", "Pre-20 realized vol", lambda v: v >= .20, "At least 20% annualized", "Below 20% annualized"),
        ("VIX at anchor", "VIX close", lambda v: v >= 25, "VIX at least 25", "VIX below 25"),
        ("Yield curve, prior observation", "10Y-2Y (pp)", lambda v: v < 0, "10Y below 2Y", "10Y at least 2Y"),
        ("CPI YoY, approx. 60-day lag", "CPI YoY", lambda v: v >= .04, "At least 4%", "Below 4%"),
    ]
    for desc,col,fn,yes,no in conditions:
        for side in [True, False]:
            label = yes if side else no
            ix = c[col].notna() & c[col].map(lambda v: fn(v)==side if pd.notna(v) else False)
            years = set(c.loc[ix,"Year"])
            for h in [20, 60, 120]:
                sub = evh[(evh.Asset=="SPX") & (evh.Horizon==h) & (evh.Year.isin(years))].copy().reset_index(drop=True)
                for q, outcome in [metrics[0], metrics[5]]:
                    observed = outcome(sub)
                    res = probability(observed,sub.Year)
                    eligible_years = ", ".join(str(int(y)) for y in sub.loc[observed.notna(),"Year"])
                    row={"Condition":desc, "Group":label, "Sessions":h, "Question":q, "Count":res["k"], "N":res["n"],
                         "Frequency":res["p"], "95% low":res["lo"], "95% high":res["hi"],
                         "Eligible years":eligible_years, "Success years":res["years"],
                         "Caution":"Too few" if res["n"]<5 else ("Very small" if res["n"]<10 else "Small")}
                    regime.append(row)
    return {"horizon_summary":pd.DataFrame(summary), "pre_summary":pd.DataFrame(pre_summary), "patterns":pd.DataFrame(patterns),
            "break_summary":pd.DataFrame(breaksum), "intraday_summary":pd.DataFrame(intra_sum),
            "all_probabilities":pd.DataFrame(allprob), "regime":pd.DataFrame(regime),
            "benchmark":pd.DataFrame(benchmark)}


def make_audit(market, events, cal, px, evh):
    rows=[]
    for asset,p in px.items():
        raw = market[f"{ASSETS[asset]}_close"].dropna()
        d = {"Asset":asset, "First close":raw.index.min(), "Last close":raw.index.max(), "Raw observations":len(raw)}
        for h in HORIZONS:
            g=evh[(evh.Asset==asset)&(evh.Horizon==h)]
            d[f"End N {h}"]=int(g["Endpoint available"].sum())
            d[f"Path N {h}"]=int(g["Complete path"].sum())
        d["Interpretation"] = ("Adjusted-minute input; adjustment and last-minute completeness not independently verified" if asset in ETFS
                               else "Price index; NASDAQ 1974–82 has close-only OHLC" if asset=="NASDAQ"
                               else "Price index; SPX-only closure dates excluded and gaps retained" if asset=="SPX"
                               else "Price index; reference session calendar")
        rows.append(d)
    extra = market[(market.index >= pd.Timestamp("1961-01-01")) & market.spx_close.notna() & market.dji_close.isna()]
    excluded = evh[(~evh["Complete path"]) & evh["Anchor close"].notna()][["Asset","Year","Horizon","Endpoint available","Exclusion detail"]].copy()
    excluded = excluded[excluded["Exclusion detail"].astype(str).str.len() > 0]
    return pd.DataFrame(rows), extra.index, excluded


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data", type=Path, default=Path(__file__).resolve().parent / "data",
                        help="Project data folder (default: data next to this script); data/processed and flat layouts are supported.")
    parser.add_argument("--out", type=Path, default=Path(__file__).resolve().parent / "Midterm_Market_Path_Study.xlsx",
                        help="Output workbook path (default: project folder next to this script).")
    args=parser.parse_args()
    print("Loading input files:")
    market, events, cal, px = load_data(args.data)
    macro = load_macro(args.data)
    context, horizons, paths, breaks, intraday, comparisons, pre_rows = [],[],[],[],[],[],[]
    for ev in events.itertuples():
        for asset in ASSETS:
            c,h,p,b,i=event_study(asset, ev.election_year, ev.election_date, "Midterm", px[asset], cal, macro, True)
            if c is not None: context.append(c)
            horizons += h; paths += p; breaks += b; intraday += i
            pre_rows += pre_election_study(asset, ev.election_year, ev.election_date, px[asset], cal)
    # One comparable November Tuesday in odd years, within the study's endpoints.
    # Odd years may have local/state elections; these are not randomized controls.
    for year in range(int(events.election_year.min()), int(events.election_year.max())+1):
        if year % 2 == 1:
            for asset in ASSETS:
                c,h,_,_,_=event_study(asset,year,nov_tuesday(year),"Odd-year November",px[asset],cal,macro,False)
                comparisons += h
    ctx, evh, path, br, intra, comp, pre = map(pd.DataFrame, [context,horizons,paths,breaks,intraday,comparisons,pre_rows])
    tables=build_summaries(ctx,evh,br,intra,comp,events,pre)
    audit, extra, excluded = make_audit(market,events,cal,px,evh)
    tables.update({"context":ctx,"pre_detail":pre,"event_horizons":evh,"paths":path,"break_detail":br,"intraday_detail":intra,
                   "comparison_detail":comp,"audit":audit,"extra_spx_dates":extra,"excluded":excluded,
                   "elections":events,"market":market,"calendar":cal})
    # create_workbook is included in this self-contained script below; there is
    # no separate workbook.py to install or copy into the project.
    args.out.parent.mkdir(parents=True,exist_ok=True)
    create_workbook(tables,args.out)
    print("Workbook:",args.out.resolve())
    print("Elections:", len(events), int(events.election_year.min()), "to", int(events.election_year.max()))
    print("Proxy-calendar SPX-only dates excluded:",len(extra))
    print(audit.to_string(index=False))
    print("\nSPX horizon summary:\n",tables["horizon_summary"].query("Asset=='SPX'").to_string(index=False))
    print("\nSPX 1974/2002 context:\n",ctx[(ctx.Asset=="SPX") & (ctx.Year.isin([1974,2002]))][["Year","Anchor date","Anchor close","Pre-20 return","Gap to 252-session peak"]].to_string(index=False))


# --- Workbook presentation, bundled here to avoid extra file dependencies ---

import datetime as dt
import math
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


NAVY = "16324A"
DARK = "0E2638"
TEAL = "087E83"
PALE = "E7F3F1"
LIGHT = "F2F6F9"
AMBER = "FFF2D7"
MUTED = "5A6B79"
WHITE = "FFFFFF"
BORDER = "D6E1E7"
GREEN = "087D66"
RED = "AC4050"
PCT = '0.0%;[Red]-0.0%;0.0%'
NUM = '#,##0.00;[Red]-#,##0.00;0.00'


def safe(value):
    if value is None:
        return None
    if isinstance(value, (float, np.floating)) and not np.isfinite(value):
        return None
    if value is pd.NaT or (not isinstance(value,(str, list,dict)) and pd.isna(value)):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime().replace(tzinfo=None)
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.bool_):
        return bool(value)
    if isinstance(value, np.floating):
        return float(value)
    return value


def wide_title(ws, title, subtitle, max_col):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=max_col)
    a=ws.cell(1,1,title)
    a.fill=PatternFill("solid",fgColor=NAVY)
    a.font=Font(name="Aptos Display",size=17,bold=True,color=WHITE)
    a.alignment=Alignment(vertical="center")
    for row in ws.iter_rows(min_row=1,max_row=1,min_col=1,max_col=max_col):
        for c in row: c.fill=PatternFill("solid",fgColor=NAVY)
    ws.row_dimensions[1].height=30
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=max_col)
    a=ws.cell(2,1,subtitle)
    a.font=Font(name="Aptos",size=10,color=MUTED,italic=True)
    a.alignment=Alignment(vertical="center",wrap_text=True)
    ws.row_dimensions[2].height=32


def table_sheet(wb, name, title, subtitle, df, pct_cols=None, width_overrides=None, freeze="C5", table_name=None):
    ws=wb.create_sheet(name)
    pct_cols=set(pct_cols or [])
    width_overrides=width_overrides or {}
    df=df.copy()
    wide_title(ws,title,subtitle,max(1,len(df.columns)))
    for j,col in enumerate(df.columns,1):
        c=ws.cell(4,j,str(col));c.font=Font(name="Aptos",bold=True,color=WHITE,size=10)
        c.fill=PatternFill("solid",fgColor=TEAL)
        c.alignment=Alignment(wrap_text=True,vertical="center")
        sample=[str(x) for x in df[col].head(150).tolist() if pd.notna(x)]
        typical=max([len(str(col))]+[min(len(x),32) for x in sample]+[9])+2
        width=max(12,min(32,typical))
        if col in {"Exclusion detail", "Interpretation", "Success years", "Question", "Condition", "Group"}:
            width=46 if col in {"Question","Condition","Exclusion detail","Interpretation"} else 30
        ws.column_dimensions[get_column_letter(j)].width=width_overrides.get(col,width)
    ws.row_dimensions[4].height=32
    for i,row in enumerate(df.itertuples(index=False,name=None),5):
        for j,(col,value) in enumerate(zip(df.columns,row),1):
            value=safe(value)
            if value is None:
                continue
            cell=ws.cell(i,j,value)
            cell.font=Font(name="Aptos",size=10,color=DARK)
            cell.alignment=Alignment(vertical="center")
            if col in pct_cols:
                cell.number_format=PCT
            elif isinstance(value,(dt.datetime,dt.date)):
                cell.number_format="yyyy-mm-dd"
            elif isinstance(value,(int,float)) and not isinstance(value,bool) and ("close" in col.lower() or "high" in col.lower() or "low" in col.lower() or "yield" in col.lower() or "peak" in col.lower()):
                cell.number_format=NUM
        if i%2==0:
            for c in ws[i]:
                c.fill=PatternFill("solid",fgColor="F7FAFB")
    if len(df) and len(df.columns):
        ref=f"A4:{get_column_letter(len(df.columns))}{4+len(df)}"
        tab=Table(displayName=table_name or ("T"+"".join(x for x in name if x.isalnum())[:24]),ref=ref)
        tab.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showFirstColumn=False,showLastColumn=False,showRowStripes=True,showColumnStripes=False)
        ws.add_table(tab)
    ws.freeze_panes=freeze
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.fitToWidth=1
    ws.print_title_rows="1:4"
    return ws


def pct_text(x, digits=1):
    return f"{x:.{digits}%}" if pd.notna(x) else "n/a"


def ratio(k,n):
    return f"{int(k)} / {int(n)}" if n else "n/a"


def dashboard(wb, t):
    ws=wb.active;ws.title="START HERE"
    ws.sheet_view.showGridLines=False
    widths={"A":21,"B":15,"C":15,"D":15,"E":20,"F":15,"G":15,"H":18,"I":3,"J":15,"K":16,"L":16,"M":16}
    for c,w in widths.items():ws.column_dimensions[c].width=w
    wide_title(ws,"MIDTERM MARKET PATHS  |  Decision brief",
               "Historical frequencies from your supplied data • 16 midterms, 1962–2022 • 5 / 10 / 20 / 30 / 60 / 90 / 120 trading sessions",8)
    def banner(row,start,end,text,color=TEAL):
        ws.merge_cells(start_row=row,start_column=start,end_row=row,end_column=end)
        for rowcells in ws.iter_rows(min_row=row,max_row=row,min_col=start,max_col=end):
            for c in rowcells:c.fill=PatternFill("solid",fgColor=color)
        c=ws.cell(row,start,text);c.font=Font(name="Aptos",bold=True,color=WHITE,size=11)
        c.alignment=Alignment(vertical="center");ws.row_dimensions[row].height=23
    def box(row,start,end,text,fill=LIGHT,size=11,bold=False,color=DARK,height=26):
        ws.merge_cells(start_row=row,start_column=start,end_row=row,end_column=end)
        for cells in ws.iter_rows(min_row=row,max_row=row,min_col=start,max_col=end):
            for c in cells:c.fill=PatternFill("solid",fgColor=fill)
        c=ws.cell(row,start,text);c.font=Font(name="Aptos",size=size,bold=bold,color=color)
        c.alignment=Alignment(vertical="center",wrap_text=True,indent=1)
        ws.row_dimensions[row].height=max(ws.row_dimensions[row].height or 0,height)
    h=t["horizon_summary"].query("Asset=='SPX'").set_index("Sessions")
    h60,h120=h.loc[60],h.loc[120]
    banner(4,1,4,"SPX | 60 sessions after the election")
    banner(4,5,8,"SPX | 120 sessions after the election")
    box(5,1,4,f"{ratio(h60['Up count'],h60['Endpoint N'])} finished higher  |  {pct_text(h60['Up frequency'])}",PALE,13,True)
    box(5,5,8,f"{ratio(h120['Up count'],h120['Endpoint N'])} finished higher  |  {pct_text(h120['Up frequency'])}",PALE,13,True)
    box(6,1,4,f"95% Wilson interval: {pct_text(h60['Up 95% low'])}–{pct_text(h60['Up 95% high'])}")
    box(6,5,8,f"95% Wilson interval: {pct_text(h120['Up 95% low'])}–{pct_text(h120['Up 95% high'])}")
    box(7,1,4,f"{ratio(h60['Dip >=10% count'],h60['Path N'])} had a close at least 10% below the anchor")
    box(7,5,8,f"{ratio(h120['Dip >=10% count'],h120['Path N'])} fell at least 10% along the way")
    banner(9,1,8,"THE PATH AFTER A BREAK  |  not just the final return")
    b=t["break_summary"].query("Asset=='SPX' and Break=='Upside' and `Forward sessions`==60").iloc[0]
    box(10,1,8,(f"First NEW upside close crossing within 60 election sessions: {ratio(b['Reentered pre-range count'],b['N'])} "
                f"returned inside the original range during the next 60 sessions."),PALE,11,True, height=30)
    box(11,1,8,(f"Of the same {int(b['N'])} observed upside-break cases, {int(b['Reentered AND finished above breakout count'])} "
                f"both re-entered the range and finished above the breakout close 60 sessions later."),LIGHT,11,False,height=30)
    box(12,1,8,"Breakouts here use 20 pre-anchor CLOSING prices. A breakout horizon begins on the breakout close, not Election Day. See After a break and Definitions.",LIGHT,10,False,height=34)
    banner(14,1,8,"TAIL RISK  |  the 1974 and 2002 cases are not interchangeable")
    ctx=t["context"].query("Asset=='SPX' and Cohort=='Midterm'")
    p=probability(ctx["Gap to 252-session peak"].map(lambda v:v<=-.20 if pd.notna(v) else np.nan),ctx.Year)
    box(15,1,8,(f"At the election anchor, {ratio(p['k'],p['n'])} SPX cases were ≥20% below their preceding 252-session peak "
                f"({p['years']}); 95% Wilson interval {pct_text(p['lo'])}–{pct_text(p['hi'])}."),AMBER,10,False,height=34)
    eh=t["event_horizons"]
    for rr,year in [(16,1974),(17,2002)]:
        c=ctx[ctx.Year==year].iloc[0]
        row=eh[(eh.Asset=="SPX")&(eh.Year==year)&(eh.Horizon==120)].iloc[0]
        box(rr,1,8,(f"{year}: pre-20 return {c['Pre-20 return']:+.1%}; at-anchor gap to prior-year peak {c['Gap to 252-session peak']:+.1%}; "
                    f"120-session result {row['Return']:+.1%}, with a worst close {row['Worst from anchor']:+.1%} from anchor."),LIGHT,10,False,height=30)
    box(18,1,8,"These describe different windows: being 20% underwater at the election is NOT the same event as a 20% decline after it. The thresholds are illustrative, not a crash forecast.",AMBER,10,False,height=34)
    banner(20,1,8,"READ THIS BEFORE SHARING")
    box(21,1,8,"Each election is one observation. Frequencies are descriptive, not predictions or causal effects; asset series and horizon windows are correlated. Confidence intervals assume an idealized binomial model.",LIGHT,10,False,height=35)
    box(22,1,8,"SPX 1994 has usable endpoints but is excluded from 60+ session path measures because three intervening closes are missing. QQQ is truncated in 2019; early NASDAQ OHLC is close-only. ETF adjustment is unverified.",LIGHT,10,False,height=38)
    box(23,1,8,"No war/conflict file was supplied. The regime and same-calendar odd-year comparisons are exploratory; they do not establish causality.",LIGHT,10,False,height=31)
    banner(25,1,8,"SPX SUMMARY  |  endpoint vs path risk")
    heads=["Sessions","Up / N","Up %","95% low","95% high","Median return","≥10% dip / N","≥10% dip %"]
    for j,head in enumerate(heads,1):
        c=ws.cell(26,j,head);c.fill=PatternFill("solid",fgColor=TEAL);c.font=Font(name="Aptos",bold=True,color=WHITE,size=10)
        c.alignment=Alignment(wrap_text=True,vertical="center")
    ws.row_dimensions[26].height=28
    for i,horizon in enumerate(HORIZONS,27):
        r=h.loc[horizon]
        vals=[horizon,ratio(r["Up count"],r["Endpoint N"]),r["Up frequency"],r["Up 95% low"],r["Up 95% high"],r["Median return"],ratio(r["Dip >=10% count"],r["Path N"]),r["Dip >=10% frequency"]]
        for j,v in enumerate(vals,1):
            c=ws.cell(i,j,safe(v));c.font=Font(name="Aptos",size=10,color=DARK)
            if j in (3,4,5,6,8):c.number_format=PCT
            if i%2==0:c.fill=PatternFill("solid",fgColor=LIGHT)
    ws.row_dimensions[34].height=30
    ws.merge_cells("A34:H34");ws["A34"]="Next: Post horizons → Break directions → After a break → Cases → Regimes → Data audit → Definitions"
    ws["A34"].font=Font(name="Aptos",bold=True,color=TEAL,size=10);ws["A34"].alignment=Alignment(wrap_text=True)
    # Numeric source for an unobtrusive, editable Excel chart.
    for j,v in enumerate(["Sessions","Finished higher","Dip >=10%"],10):ws.cell(25,j,v)
    for i,horizon in enumerate(HORIZONS,26):
        r=h.loc[horizon];ws.cell(i,10,horizon);ws.cell(i,11,float(r["Up frequency"]));ws.cell(i,12,float(r["Dip >=10% frequency"]))
    chart=BarChart();chart.type="col";chart.title="SPX: endpoint and path risk";chart.y_axis.title="Historical frequency";chart.x_axis.title="Sessions"
    chart.style=10;chart.height=8;chart.width=18
    chart.add_data(Reference(ws,min_col=11,max_col=12,min_row=25,max_row=32),titles_from_data=True)
    chart.set_categories(Reference(ws,min_col=10,min_row=26,max_row=32))
    chart.y_axis.numFmt="0%";chart.y_axis.scaling.min=0;chart.y_axis.scaling.max=1
    ws.add_chart(chart,"A36")
    ws.freeze_panes="A4";ws.sheet_view.zoomScale=90
    return ws


def make_cases(t):
    ctx=t["context"][(t["context"].Asset=="SPX") & (t["context"].Year.isin([1974,2002,2018]))]
    eh=t["event_horizons"][(t["event_horizons"].Asset=="SPX") & (t["event_horizons"].Year.isin([1974,2002,2018]))].copy()
    keep=ctx[["Year","Pre-20 return","Gap to 252-session peak"]]
    d=eh.merge(keep,on="Year",how="left")
    return d[["Year","Horizon","Anchor date","End date","Pre-20 return","Gap to 252-session peak","Anchor state",
              "Return","Worst from anchor","Max drawdown","First new cross","First cross session"]].sort_values(["Year","Horizon"])


def crash_risk_table(t):
    """Make pre-election bear state and post-election losses visibly distinct."""
    rows=[]
    for asset in ["SPX","DJI","NASDAQ","SPY","QQQ"]:
        c=t["context"][(t["context"].Asset==asset) & (t["context"].Cohort=="Midterm")]
        p=probability(c["Gap to 252-session peak"].map(lambda x: x<=-.20 if pd.notna(x) else np.nan), c.Year)
        rows.append({"Asset":asset,"Window":"At anchor","Sessions":0,
                     "Question":"Anchor >=20% below prior 252-session closing peak",
                     "Count":p["k"],"N":p["n"],"Frequency":p["p"],"95% low":p["lo"],"95% high":p["hi"],"Success years":p["years"]})
    wanted=["Fall >=10% from anchor at any close", "Fall >=20% from anchor at any close", "Peak-to-trough drawdown >=10%"]
    a=t["all_probabilities"]
    subset=a[(a.Section=="Election horizon") & (a.Question.isin(wanted))]
    for _,r in subset.iterrows():
        rows.append({"Asset":r["Asset"],"Window":"After anchor","Sessions":r["Horizon"],"Question":r["Question"],
                     "Count":r["Count"],"N":r["N"],"Frequency":r["Frequency"],"95% low":r["95% low"],
                     "95% high":r["95% high"],"Success years":r["Success years"]})
    order={a:i for i,a in enumerate(["SPX","DJI","NASDAQ","SPY","QQQ"])}
    out=pd.DataFrame(rows);out["_order"]=out.Asset.map(order)
    return out.sort_values(["_order","Sessions","Question"]).drop(columns="_order").reset_index(drop=True)


def add_case_charts(wb, case_ws, paths):
    ws=wb.create_sheet("Chart data")
    ws.sheet_state="hidden"
    ws.cell(1,1,"Session")
    for i,day in enumerate(range(-20,121),2):ws.cell(i,1,day)
    colors=["087E83","D48E24","A6404F"]
    for idx,year in enumerate([1974,2002,2018]):
        subset=paths[(paths.Asset=="SPX")&(paths.Year==year)].set_index("Session")
        start=2+idx*4
        for off,label in enumerate([f"{year} close / anchor", "Pre-20 high / anchor", "Pre-20 low / anchor", "Anchor"]):ws.cell(1,start+off,label)
        for row,day in enumerate(range(-20,121),2):
            if day in subset.index:
                item=subset.loc[day]
                for off,key in enumerate(["Close / anchor","Pre-high / anchor","Pre-low / anchor"]):
                    ws.cell(row,start+off,safe(item[key]))
                ws.cell(row,start+3,1.0)
        chart=LineChart();chart.title=f"SPX {year}: close path & pre-election range";chart.y_axis.title="Anchor close = 100%";chart.x_axis.title="Trading session relative to anchor"
        chart.height=7.3;chart.width=20;chart.style=13;chart.legend.position="b"
        chart.visible_cells_only=False  # source sheet is hidden; still plot it in Excel
        chart.add_data(Reference(ws,min_col=start,max_col=start+3,min_row=1,max_row=142),titles_from_data=True)
        chart.set_categories(Reference(ws,min_col=1,min_row=2,max_row=142));chart.y_axis.numFmt="0%"
        for j,series in enumerate(chart.series):
            series.graphicalProperties.line.solidFill=[colors[idx],"D28B22","8497A5","263A4A"][j]
            series.graphicalProperties.line.width=28000 if j==0 else 15000
            if j>0:series.graphicalProperties.line.dashStyle="sysDash"
        case_ws.add_chart(chart,f"N{5+idx*16}")


def definitions_sheet(wb,t):
    defs=[
        ("Purpose", "Translate the supplied historical midterm sample into endpoint, path, breakout and risk frequencies. No model forecasts a future election."),
        ("Sample", "16 supplied midterm events, 1962–2022. Assets are analyzed separately. Horizons are 5, 10, 20, 30, 60, 90 and 120 trading sessions."),
        ("D0 / anchor", "Last DJI-calendar trading session on or before Election Day; close is the anchor. 1962 and 1966 Election Day were market-closed in the supplied calendar. Election results were generally not known at the anchor close."),
        ("Session calendar", "The supplied DJI daily-close dates proxy the shared exchange calendar. SPX-only dates when DJI is absent are excluded; no formal independent exchange-calendar audit has been performed."),
        ("Pre-election window", "A pre-h-session return is anchor close / close h sessions earlier minus 1. The prior-20 range uses the 20 sessions strictly BEFORE the anchor (not including D0), matching the notebook's window boundary."),
        ("Prior-20 CLOSE range", "Primary pre-high = maximum of those 20 daily CLOSES; pre-low = minimum of those CLOSES. This is not the notebook's pre-event intraday high and low. The Intraday touches tab separately uses true prior-bar high and low where data permit."),
        ("Above / below pre-range", "Any post-anchor close strictly above the pre-high or below the pre-low, regardless of whether the anchor was already outside. High and low counts can overlap."),
        ("NEW close crossing", "An upside new cross requires previous close <= pre-high and today's close > pre-high. A downside cross requires previous close >= pre-low and today's close < pre-low. Already outside on D0 is not a new crossing; a later re-entry and recross can qualify."),
        ("First new cross", "Earliest new upside or downside closing cross after D0. Upside, downside and none form an exclusive partition for each complete window. None does not mean the anchor was inside the range."),
        ("Follow-through cohort", "First NEW upside crossing and first NEW downside crossing are analyzed separately, regardless of which side crossed first. Search is limited to the first 60 post-election sessions and requires that search window complete. Forward horizons then start at that side's crossing close and must be fully observed."),
        ("Re-entry", "At least one subsequent close between the original pre-20 closing low and high, inclusive. A close-confirmed breakout is already outside the level; re-entry may be temporary."),
        ("Pullback after a break", "A close at or below 97% of the breakout close during the specified forward period. Pullback-and-finish-higher requires that event and a positive return from breakout close to the forward-horizon close."),
        ("Directional continuation", "At the forward horizon, >=+3% from an upside breakout close or <=-3% from a downside breakout close. These 3% cutoffs are conventions, not statistically optimized or universal."),
        ("Tight range / consolidation", "All closes including the initial close have max(close)/min(close)-1 <=5% in the specified window. This is one deliberately mechanical definition; longer windows are less likely to pass. A flat endpoint alone is not consolidation."),
        ("Worst dip from anchor", "Minimum of anchor and subsequent closes divided by the anchor close minus 1. Dip >=10% means a close at or below 90% of the anchor at any time, NOT a peak-to-trough drawdown."),
        ("Maximum drawdown", "Minimum close / running maximum close - 1 within the specified window, including its starting close. This differs from the worst decline relative to the original anchor."),
        ("252-session peak gap", "Anchor close / maximum close across the 252 sessions ending at the anchor minus 1. The 20%-underwater regime is an illustrative at-election state, NOT the probability of a 20% post-election crash."),
        ("Intraday touch", "Previous 20 bars: pre-high=max(high), pre-low=min(low); after anchor, any daily high above pre-high or daily low below pre-low. Require every relevant bar to have non-flat, internally consistent OHLC. If both are touched in the same bar, their ordering is unknown. This avoids treating early close-only NASDAQ bars as intraday highs/lows."),
        ("Incomplete sessions", "Endpoint return requires a valid anchor and the exact horizon-end close; intervening gaps do not invalidate the endpoint. Path metrics require every intervening DJI-calendar close. For ETFs, >=180 supplied minute bars is only a coarse screen; it does not prove the final minute was captured."),
        ("Missing windows", "Never drop gaps and compress trading time. An incomplete path is missing, not a failed event. SPX 1994 has missing 1994-12-30, 1995-01-19 and 1995-01-20 closes: eligible endpoints may remain usable, but not path statistics. QQQ ends 2019-04-18 with only 51 minutes that day."),
        ("Frequency and Wilson 95%", "Count / number of eligible, fully observed elections; Wilson score interval. The interval assumes a stable binomial process and approximate independent events—strong assumptions for regime-changing markets. It is not a forecast interval."),
        ("N = 0", "No eligible observations: count is 0, frequency and interval are blank. The count 0 is NOT an observed 0% probability when N=0."),
        ("Odd-year comparator", "Same first-Tuesday-after-first-Monday in November for odd years between 1962 and 2022. These years are not federal presidential/midterm general-election years but may have local/state elections. Matching season does not control macro regime, politics or secular trends; no causality."),
        ("Regime comparisons", "SPX conditions are pre-anchor return, 252-session peak gap, 20-session realized volatility, VIX, prior 10Y-2Y and CPI YoY. One factor at a time, thresholds explicit, 20/60/120 sessions; tiny subgroups and many comparisons are exploratory, not reliable causal effects."),
        ("VIX and rates", "VIX close on/before anchor. Fed funds and yields use observations strictly before anchor and no older than 7 calendar days; yield spread requires the 2Y and 10Y observation dates to match. Publication timing and revisions were not reconstructed."),
        ("CPI timing", "Latest CPI observation whose observation-month date plus 60 calendar days is on/before the anchor; year-over-year growth from the supplied series. This lag is a proxy, NOT verified release dates or a real-time vintage; the input may contain revisions."),
        ("ETF adjustment", "SPY/QQQ OHLC is taken exactly from the provided market_daily.csv, described in the notebook as aggregated adjusted minute files. The adjustment methodology and complete dividend reinvestment cannot be verified from the supplied files. Do not call the ETF numbers proven total returns."),
        ("Trading and execution", "Index returns are price-index changes, not investable total returns. No transaction costs, slippage, taxes, leverage or realistic order execution modeled; close signals cannot automatically be executed at the same closing price."),
        ("Conflicts and causality", "No war/conflict data file was supplied. Recession flags or historical political results were also not supplied in usable form. We do not infer missing categories, causal effects or statistical significance from these associations."),
        ("Multiple comparisons", "Across assets, horizons, thresholds and conditions, many patterns are inspected on the same events. No multiple-testing correction or out-of-sample validation is performed. The most attractive result can be a selection artifact."),
    ]
    df=pd.DataFrame(defs,columns=["Term","Definition and interpretation"])
    ws=table_sheet(wb,"Definitions","Definitions & limitations",
                   "Read these before using any frequency. Values are static calculations; rerun the script to change thresholds or update data.",df,
                   width_overrides={"Term":30,"Definition and interpretation":115},freeze="B5")
    for row in range(5,5+len(df)):
        ws.row_dimensions[row].height=48 if row not in (5,6) else 34
        ws.cell(row,2).alignment=Alignment(wrap_text=True,vertical="center")
    return ws


def create_workbook(t, output_path: Path):
    wb=Workbook();wb.properties.title="Midterm market paths | 1962–2022"
    wb.properties.subject="Descriptive election-event frequencies and path risks"
    wb.properties.creator="Arena research assistant"
    dashboard(wb,t)
    hs=t["horizon_summary"]
    table_sheet(wb,"Post horizons","Post-election endpoints & risk",
        "All returns and dips use closing prices. Endpoint N requires only D0 and horizon-end closes; path N also requires every intervening close. Wilson intervals appear in All probabilities.",hs,
        pct_cols=["Up frequency","Up 95% low","Up 95% high","Median return","Lowest return","Median worst dip",
                  "Dip >=5% frequency","Dip >=10% frequency","Dip >=20% frequency","Tight frequency"])
    table_sheet(wb,"Break directions","Closing-range break directions",
        "First NEW up / down / no crossing are mutually exclusive; any above / below can overlap and may already be true at the anchor. N requires a complete pre-range and post window.",t["patterns"],
        pct_cols=[c for c in t["patterns"].columns if "frequency" in c.lower()])
    table_sheet(wb,"After a break","What happens after a NEW closing breakout?",
        "Condition: first side-specific NEW crossing within 60 post-election sessions; forward windows start at the breakout close. Rows with missing follow-through are excluded, not failures.",t["break_summary"],
        pct_cols=[c for c in t["break_summary"].columns if "frequency" in c.lower()])
    cases=make_cases(t)
    case_ws=table_sheet(wb,"Cases","Case studies  |  SPX 1974, 2002 and 2018",
        "Compare return with the deepest closing loss from the same election anchor. The plotted prior range uses pre-anchor closes; 2018 is included as another >=10% post-anchor dip.",cases,
        pct_cols=["Pre-20 return","Gap to 252-session peak","Return","Worst from anchor","Max drawdown"],
        width_overrides={"Anchor state":20})
    add_case_charts(wb,case_ws,t["paths"])
    table_sheet(wb,"Crash risk","Large losses  |  before vs after the election",
        "At-anchor 252-session peak gap and after-anchor declines are different events. A zero count is not a zero future probability; 95% Wilson intervals are illustrative.",crash_risk_table(t),
        pct_cols=["Frequency","95% low","95% high"],width_overrides={"Question":55,"Success years":45})
    table_sheet(wb,"Pre-election","What happened BEFORE the election anchor?",
        "Return from h sessions before D0 through D0. The 20-session pre-close range used for breakout thresholds excludes D0; see Definitions.",t["pre_summary"],
        pct_cols=["Up frequency","95% low","95% high","Median return","Median max drawdown"])
    table_sheet(wb,"Regimes","Exploratory conditional frequencies  |  SPX",
        "Single-factor conditions and two outcomes at 20 / 60 / 120 sessions. Tiny N, revised data and multiple comparisons: association only, not causation or a verified trading signal.",t["regime"],
        pct_cols=["Frequency","95% low","95% high"],width_overrides={"Condition":35,"Group":32,"Question":40,"Eligible years":45,"Success years":45})
    table_sheet(wb,"Calendar comparator","Seasonal comparison  |  odd-year November",
        "Odd-year first Tuesday after the first Monday in November, 1963–2021. Some have local/state elections. Difference is percentage points, descriptive and not a causal estimate.",t["benchmark"],
        pct_cols=["Midterm frequency","Midterm 95% low","Midterm 95% high","Odd-year frequency","Odd-year 95% low","Odd-year 95% high"],
        width_overrides={"Question":42})
    table_sheet(wb,"Intraday touches","Intraday levels  |  separate from closing-range breaks",
        "Prior 20 sessions: max(high), min(low). A touch uses high/low; high and low can both occur. Invalid, flat or close-only OHLC windows are excluded. Same-day order is unknown.",t["intraday_summary"],
        pct_cols=[c for c in t["intraday_summary"].columns if "frequency" in c.lower()])
    table_sheet(wb,"All probabilities","All probability estimates & Wilson intervals",
        "Filter by section, asset, horizon and question. A blank frequency with N=0 means no observations—not 0%. Success years support a case-by-case audit.",t["all_probabilities"],
        pct_cols=["Frequency","95% low","95% high"],width_overrides={"Section":25,"Condition":38,"Question":49,"Success years":50})
    ev=t["event_horizons"].drop(columns=["Cohort"])
    table_sheet(wb,"Event ledger","Election-by-election horizon ledger",
        "One row per asset × midterm × horizon. Endpoint return may remain available even if the path is incomplete. Inspect Exclusion detail; first cross means a NEW closing-price crossing.",ev,
        pct_cols=["Return","Worst from anchor","Max drawdown","Close range"],width_overrides={"Exclusion detail":65,"Anchor state":21})
    table_sheet(wb,"Breakout detail","Breakout follow-through  |  case-level audit",
        "Only observed first side-specific new crosses in a complete 60-session search are listed. An incomplete forward horizon has blank outcomes; break dates are closing-price signals.",t["break_detail"],
        pct_cols=["Forward return","Worst after break"],width_overrides={"Anchor state":20})
    table_sheet(wb,"Intraday detail","Intraday touches  |  case-level audit",
        "An ineligible row has insufficient or flat OHLC. If both levels are hit on the first touch day, the intraday sequence is unknown.",t["intraday_detail"],
        width_overrides={"First intraday touch":36})
    ctx=t["context"].drop(columns=["Cohort"])
    table_sheet(wb,"Event context","Election anchor & pre-event context",
        "252-session peak and volatility require complete pre-event closes. Rates are prior observations; CPI uses an approximate lag and revised data; VIX is only available from 1990.",ctx,
        pct_cols=["Pre-20 return","Gap to 252-session peak","Pre-20 realized vol","CPI YoY"],width_overrides={"Anchor state":20})
    table_sheet(wb,"Daily paths","Daily event paths  |  normalized to anchor",
        "One row per valid anchor and session -20 to +120 on the DJI calendar. Close/anchor = 100% at D0; blank/missing values are not filled. ETF minute bars below 180 are unusable.",t["paths"],
        pct_cols=["Close / anchor","Pre-high / anchor","Pre-low / anchor"],freeze="D5")
    table_sheet(wb,"Data audit","Data coverage & usable event counts",
        "End N = valid endpoint returns; Path N = fully observed paths. The two can differ. Coverage is not the same as an asset's inception date.",t["audit"],
        width_overrides={"Interpretation":68})
    issues=[]
    issues.append({"Issue":"SPX extra dates outside DJI calendar","Evidence":f"{len(t['extra_spx_dates'])} SPX-only dates since 1961, including market-closed holidays in 1982–86 and 2001-09-11.","Treatment":"Excluded from the common DJI-based session calendar; listed below."})
    issues.append({"Issue":"SPX 1994 incomplete future","Evidence":"Missing closes on 1994-12-30, 1995-01-19 and 1995-01-20 on the DJI calendar.","Treatment":"60 / 90 / 120 session path metrics and breakout search exclude 1994; actual horizon-end returns are retained if both endpoints exist."})
    issues.append({"Issue":"QQQ truncated","Evidence":"Last supplied close 2019-04-18; that session has only 51 minute bars. There is no QQQ data in 2022.","Treatment":"2018 QQQ 120-session window excluded; no imputation. 2022 QQQ unavailable."})
    issues.append({"Issue":"Early NASDAQ OHLC","Evidence":"1974, 1978 and 1982 event windows have high=low=close throughout.","Treatment":"Closing-price metrics retained; flat bars are ineligible for intraday-touch statistics."})
    issues.append({"Issue":"ETF minute bars / adjustments","Evidence":"The notebook labels raw input adjusted minute files, but methodology and final-minute coverage cannot be verified from daily aggregation.","Treatment":"Require >=180 bars; historical ETF returns are not asserted to be verified total returns."})
    for d in t["extra_spx_dates"]:
        issues.append({"Issue":"Excluded SPX-only date", "Evidence":str(d.date()), "Treatment":"Not in DJI-based reference calendar."})
    table_sheet(wb,"Data issues","Quality issues & excluded calendar dates",
        "The DJI daily series is a proxy calendar, not a separately authenticated exchange calendar. Review the exact dates and exclusions before presenting results.",pd.DataFrame(issues),
        width_overrides={"Issue":34,"Evidence":80,"Treatment":80},freeze="B5")
    table_sheet(wb,"Excluded windows","Incomplete post-election windows",
        "Paths with a valid anchor rejected for gaps, short ETF sessions or calendar end. Some exact endpoint returns remain available. Assets with no anchor appear in Event ledger.",t["excluded"],
        width_overrides={"Exclusion detail":100},freeze="B5")
    definitions_sheet(wb,t)
    wb.active=0
    output_path=Path(output_path)
    wb.save(output_path)


if __name__ == "__main__":
    main()
