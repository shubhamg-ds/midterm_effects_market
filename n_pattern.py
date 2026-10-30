import numpy as np
import pandas as pd
from pathlib import Path
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# ============================ SETTINGS ============================
FOLDER = Path(r"C:\projects\research\mid_terms_effects\data")                       # folder with the CSVs
FILES = {"SPX": "SPX_1D.csv", "DJI": "DJI_1D.csv", "NASDAQ": "NASDAQ_DLY_IXIC_1D.csv"}
OUT = "midterm_study.xlsx"

PRE_WINDOW = 60      # calendar days before election searched for the auto-detected low
RETEST_DAYS = 60     # calendar days after election searched for the retest low (~to early Jan)
NEAR_TOL = 3.0       # % : retest within +3% above the old low (or anything below) counts as "N retest"
MIN_RALLY = 2.0      # % : minimum low->election rally for an N to count
PATH_PRE, PATH_POST = 30, 130   # trading-day window for the average path

# Your hand-picked pre-election lows. 1998 left out on purpose -> auto-detected.
USER_LOWS = {
    1962: "1962-10-24", 1966: "1966-10-10", 1970: "1970-09-24", 1974: "1974-10-04",
    1978: "1978-10-30", 1982: "1982-10-26", 1986: "1986-09-29", 1990: "1990-10-11",
    1994: "1994-10-05", 2002: "2002-10-10", 2006: "2006-11-03", 2010: "2010-10-19",
    2014: "2014-10-21", 2018: "2018-10-29", 2022: "2022-10-13", 2026: "2026-10-23",
}
USE_USER_LOWS = True   # False -> use auto-detected lows for every year (a good robustness check)
MIDTERMS = list(range(1962, 2027, 4))
# ==================================================================


def election_date(year):
    d = pd.Timestamp(year, 11, 1)
    first_monday = d + pd.Timedelta(days=(0 - d.weekday()) % 7)
    return first_monday + pd.Timedelta(days=1)


def load(path):
    df = pd.read_csv(path)
    df.columns = [c.strip().lower() for c in df.columns]
    tcol = next(c for c in df.columns if c in ("time", "date", "datetime", "timestamp"))
    t = df[tcol]
    if pd.api.types.is_numeric_dtype(t):
        dt = pd.to_datetime(t, unit="s" if abs(t.iloc[-1]) < 1e11 else "ms")
    else:
        dt = pd.to_datetime(t.astype(str).str.strip(), format="mixed", dayfirst=True,
                            utc=True, errors="coerce").dt.tz_localize(None)
        if dt.isna().any():
            bad = t[dt.isna()].head(10).tolist()
            raise ValueError(f"{path}: {int(dt.isna().sum())} dates failed to parse, e.g. {bad}")

    df.index = pd.DatetimeIndex(dt).normalize()
    for c in ("open", "high", "low", "close"):
        if c not in df.columns:
            df[c] = df["close"]
    df = df[["open", "high", "low", "close"]].astype(float)
    df = df[~df.index.duplicated(keep="last")].sort_index().dropna(subset=["close"])
    return df


def close_asof(df, d):
    if d > df.index[-1]:
        return np.nan, pd.NaT
    p = df.index.searchsorted(d, side="right") - 1
    return df["close"].iloc[p], df.index[p]


def pct(a, b):
    return round((a / b - 1) * 100, 2) if pd.notna(a) and pd.notna(b) and b else np.nan


def analyze(df, year):
    e = election_date(year)
    idx, last = df.index, df.index[-1]
    if idx[0] > e - pd.Timedelta(days=PRE_WINDOW + 5) or last < e - pd.Timedelta(days=5):
        return None
    pre = df.loc[e - pd.Timedelta(days=PRE_WINDOW): e - pd.Timedelta(days=1)]
    if pre.empty:
        return None
    auto_low_date = pre["low"].idxmin()
    if USE_USER_LOWS and year in USER_LOWS:
        ld = idx[min(idx.searchsorted(pd.Timestamp(USER_LOWS[year])), len(idx) - 1)]
        src = "user"
    else:
        ld, src = auto_low_date, "auto"
    pl = idx.get_loc(ld)
    low_px = df["low"].iloc[pl]
    reached = last >= e
    rp = idx.searchsorted(e, side="right") - 1            # last bar on/before election
    ec = df["close"].iloc[rp]
    seg = df.iloc[pl:rp + 1]
    pk_date, pk = seg["high"].idxmax(), seg["high"].max()
    td_pre = rp - pl

    r = {"Year": year, "Election": e.date(), "Status": "complete" if reached else "in progress (to date)",
         "Low source": src, "Pre-low date": ld.date(), "Auto-low date": auto_low_date.date(),
         "Pre-low price": round(low_px, 2),
         "Low: cal days before elec": (e - ld).days, "Low: trading days before elec": td_pre,
         "Rally low->elec close %": pct(ec, low_px),
         "Rally slope %/tday": round(pct(ec, low_px) / td_pre, 3) if td_pre else np.nan,
         "N-height: peak vs low %": pct(pk, low_px),
         "Peak date": pk_date.date(), "Pullback peak->elec close %": pct(ec, pk)}

    nan_keys = ["Retest low date", "Retest low price", "Retest vs pre-low %", "Retest drop from elec close %",
                "Retest days after elec", "N width (tdays low->retest)", "N width (cal days)",
                "Ret 1d %", "Ret 5d %", "Ret 10d %", "Ret 1m %", "Ret 3m %", "Ret 6m %", "Ret 12m %",
                "Pre-low -> 6m %", "Retest low -> 6m %", "Retest->6m slope %/tday",
                "6m max drawdown from elec %", "6m worst low vs elec close %", "6m worst low date",
                "JanFeb low date", "JanFeb low vs elec close %", "JanFeb low vs pre-low %"]
    for k in nan_keys:
        r[k] = np.nan
    if not reached:
        return r

    pos_e = rp
    for n in (1, 5, 10):
        if pos_e + n < len(df):
            r[f"Ret {n}d %"] = pct(df["close"].iloc[pos_e + n], ec)
    for m in (1, 3, 6, 12):
        px, _ = close_asof(df, e + pd.DateOffset(months=m))
        r[f"Ret {m}m %"] = pct(px, ec)
    px6, d6 = close_asof(df, e + pd.DateOffset(months=6))
    r["Pre-low -> 6m %"] = pct(px6, low_px)

    if last >= e + pd.Timedelta(days=RETEST_DAYS):
        w = df.loc[e + pd.Timedelta(days=1): e + pd.Timedelta(days=RETEST_DAYS)]
        rt_date, rt_px = w["low"].idxmin(), w["low"].min()
        r["Retest low date"] = rt_date.date()
        r["Retest low price"] = round(rt_px, 2)
        r["Retest vs pre-low %"] = pct(rt_px, low_px)
        r["Retest drop from elec close %"] = pct(rt_px, ec)
        r["Retest days after elec"] = (rt_date - e).days
        r["N width (tdays low->retest)"] = idx.get_loc(rt_date) - pl
        r["N width (cal days)"] = (rt_date - ld).days
        if pd.notna(px6):
            r["Retest low -> 6m %"] = pct(px6, rt_px)
            td = idx.get_loc(d6) - idx.get_loc(rt_date)
            r["Retest->6m slope %/tday"] = round(r["Retest low -> 6m %"] / td, 3) if td > 0 else np.nan

    if pd.notna(px6):
        w6 = df.loc[e:d6, "close"]
        r["6m max drawdown from elec %"] = round(((w6 / w6.cummax() - 1).min()) * 100, 2)
        wl = df.loc[e:d6, "low"]
        r["6m worst low vs elec close %"] = pct(wl.min(), ec)
        r["6m worst low date"] = wl.idxmin().date()
    jf = df.loc[pd.Timestamp(year + 1, 1, 1): pd.Timestamp(year + 1, 3, 1) - pd.Timedelta(days=1)]
    if len(jf) and last >= jf.index[-1]:
        r["JanFeb low date"] = jf["low"].idxmin().date()
        r["JanFeb low vs elec close %"] = pct(jf["low"].min(), ec)
        r["JanFeb low vs pre-low %"] = pct(jf["low"].min(), low_px)
    return r


def add_shape(ev):
    def shape(row):
        if pd.isna(row["Retest vs pre-low %"]):
            return np.nan
        if row["Rally low->elec close %"] < MIN_RALLY:
            return "No pre-election rally"
        v = row["Retest vs pre-low %"]
        if v < -NEAR_TOL:
            return "N (retest makes lower low)"
        if v <= NEAR_TOL:
            return "N (retest ~ old low)"
        return "Higher low (no real retest)"
    ev["Shape"] = ev.apply(shape, axis=1)
    return ev


def flag(s, f):
    return f(s).astype(float).where(s.notna())


def prob_rows(ev):
    h = ev[ev["Status"] == "complete"]
    out = []

    def P(label, cond):
        c = cond.dropna()
        n, k = len(c), int(c.sum())
        out.append({"Observation": label, "Hits": k, "N": n,
                    "Probability %": round(100 * k / n, 1) if n else np.nan})

    P("Rally low->election close is UP", flag(h["Rally low->elec close %"], lambda s: s > 0))
    P("Rally low->election >= 3%", flag(h["Rally low->elec close %"], lambda s: s >= 3))
    P("Rally low->election >= 5%", flag(h["Rally low->elec close %"], lambda s: s >= 5))
    P("Day after election is UP", flag(h["Ret 1d %"], lambda s: s > 0))
    P("5 sessions after election UP", flag(h["Ret 5d %"], lambda s: s > 0))
    P("10 sessions after election UP", flag(h["Ret 10d %"], lambda s: s > 0))
    for m in (1, 3, 6, 12):
        P(f"{m}m after election UP", flag(h[f"Ret {m}m %"], lambda s: s > 0))
    P("6m return >= +5%", flag(h["Ret 6m %"], lambda s: s >= 5))
    P("6m return >= +10%", flag(h["Ret 6m %"], lambda s: s >= 10))
    P("Pre-low -> 6m UP", flag(h["Pre-low -> 6m %"], lambda s: s > 0))
    P("Retest ~ old low or lower (N retest, any depth)", flag(h["Retest vs pre-low %"], lambda s: s <= NEAR_TOL))
    P(f"Retest within +/-{NEAR_TOL}% of old low (tight N)", flag(h["Retest vs pre-low %"], lambda s: s.abs() <= NEAR_TOL))
    P("Retest makes a LOWER low (< -tol)", flag(h["Retest vs pre-low %"], lambda s: s < -NEAR_TOL))
    P("Retest holds ABOVE old low (> +tol)", flag(h["Retest vs pre-low %"], lambda s: s > NEAR_TOL))
    P("Drop from election close >= 3% within 60d", flag(h["Retest drop from elec close %"], lambda s: s <= -3))
    P("Drop from election close >= 5% within 60d", flag(h["Retest drop from elec close %"], lambda s: s <= -5))
    P("Drop from election close >= 8% within 60d", flag(h["Retest drop from elec close %"], lambda s: s <= -8))
    P("Retest low within 30 cal days after election", flag(h["Retest days after elec"], lambda s: s <= 30))
    P("Retest low -> 6m is UP", flag(h["Retest low -> 6m %"], lambda s: s > 0))
    P("Retest low -> 6m >= +10%", flag(h["Retest low -> 6m %"], lambda s: s >= 10))
    P("Worst low in 6m is below election close", flag(h["6m worst low vs elec close %"], lambda s: s < 0))
    P("Worst low in 6m is >= 5% below election close", flag(h["6m worst low vs elec close %"], lambda s: s <= -5))
    P("Jan-Feb low below election close", flag(h["JanFeb low vs elec close %"], lambda s: s < 0))
    P("Jan-Feb low below the pre-election low", flag(h["JanFeb low vs pre-low %"], lambda s: s < 0))
    P("Full N: rally>=MIN, retest<=tol, then retest->6m UP",
      ((h["Rally low->elec close %"] >= MIN_RALLY) & (h["Retest vs pre-low %"] <= NEAR_TOL) &
       (h["Retest low -> 6m %"] > 0)).astype(float).where(h["Retest low -> 6m %"].notna()))
    return pd.DataFrame(out)


def cond_rows(ev):
    h = ev[ev["Status"] == "complete"].copy()
    out = []

    def CP(cond_label, cond, outcome_label, outcome):
        m = cond.notna() & outcome.notna()
        t, f = m & (cond == 1), m & (cond == 0)
        pt = round(100 * outcome[t].mean(), 1) if t.sum() else np.nan
        pf = round(100 * outcome[f].mean(), 1) if f.sum() else np.nan
        out.append({"IF": cond_label, "THEN P(outcome)": outcome_label,
                    "N if true": int(t.sum()), "P % if true": pt,
                    "N if false": int(f.sum()), "P % if false": pf,
                    "Base P %": round(100 * outcome[m].mean(), 1) if m.sum() else np.nan})

    rally = h["Rally low->elec close %"]
    slope = h["Rally slope %/tday"]
    up6 = flag(h["Ret 6m %"], lambda s: s > 0)
    up1 = flag(h["Ret 1m %"], lambda s: s > 0)
    retest_ok = flag(h["Retest vs pre-low %"], lambda s: s <= NEAR_TOL)
    retest_lower = flag(h["Retest vs pre-low %"], lambda s: s < -NEAR_TOL)
    drop5 = flag(h["Retest drop from elec close %"], lambda s: s <= -5)
    big = flag(rally, lambda s: s > rally.median())
    steep = flag(slope, lambda s: s > slope.median())
    early = flag(h["Low: trading days before elec"], lambda s: s > h["Low: trading days before elec"].median())
    wide = flag(h["N width (tdays low->retest)"], lambda s: s > h["N width (tdays low->retest)"].median())
    first1 = flag(h["Ret 1d %"], lambda s: s > 0)
    first5 = flag(h["Ret 5d %"], lambda s: s > 0)

    CP("Rally into election bigger than median", big, "6m after election UP", up6)
    CP("Rally into election bigger than median", big, "Retest at/below old low", retest_ok)
    CP("Rally into election bigger than median", big, "Drop >=5% after election", drop5)
    CP("Rally slope steeper than median", steep, "Drop >=5% after election", drop5)
    CP("Rally slope steeper than median", steep, "Retest makes lower low", retest_lower)
    CP("Low was earlier than median (more days before elec)", early, "Retest at/below old low", retest_ok)
    CP("Low was earlier than median (more days before elec)", early, "6m after election UP", up6)
    CP("Retest at/below old low (N formed)", retest_ok, "6m after election UP", up6)
    CP("Retest makes lower low", retest_lower, "6m after election UP", up6)
    CP("Drop >=5% after election", drop5, "6m after election UP", up6)
    CP("N is wider than median", wide, "6m after election UP", up6)
    CP("Day after election UP", first1, "1m after election UP", up1)
    CP("Day after election UP", first1, "6m after election UP", up6)
    CP("First 5 sessions UP", first5, "6m after election UP", up6)
    CP("First 5 sessions UP", first5, "Retest at/below old low", retest_ok)
    return pd.DataFrame(out)


def describe(ev):
    h = ev[ev["Status"] == "complete"]
    cols = [c for c in h.columns if h[c].dtype.kind in "fi" and c != "Year"]
    d = h[cols].agg(["count", "mean", "median", "min", "max"]).T.round(2)
    d.insert(0, "Metric", d.index)
    return d.reset_index(drop=True)


def cycle_table(df):
    rows = []
    for y in range(df.index[0].year, df.index[-1].year + 1):
        e = election_date(y)
        if e > df.index[-1] or e < df.index[0]:
            continue
        ec, _ = close_asof(df, e)
        row = {"Year": y, "Cycle": {2: "Midterm", 0: "Presidential", 1: "Post-election", 3: "Pre-election"}[y % 4]}
        for m in (1, 3, 6, 12):
            px, _ = close_asof(df, e + pd.DateOffset(months=m))
            row[f"{m}m"] = pct(px, ec)
        rows.append(row)
    t = pd.DataFrame(rows)
    g = t.groupby("Cycle").agg(Years=("Year", "count"),
                               **{f"Mean {m} %": (m, "mean") for m in ("1m", "3m", "6m", "12m")},
                               **{f"Median {m} %": (m, "median") for m in ("1m", "6m")},
                               **{f"P(up) {m} %": (m, lambda s: 100 * (s.dropna() > 0).mean()) for m in ("1m", "6m")})
    return g.round(1).reset_index()


def analog_table(ev):
    cur = ev[ev["Status"] != "complete"]
    if cur.empty:
        return None, None
    c = cur.iloc[0]
    h = ev[ev["Status"] == "complete"].copy()
    h["Dist"] = ((h["Rally low->elec close %"] - c["Rally low->elec close %"]).abs() / h["Rally low->elec close %"].std() +
                 (h["Rally slope %/tday"] - c["Rally slope %/tday"]).abs() / h["Rally slope %/tday"].std())
    cols = ["Year", "Rally low->elec close %", "Rally slope %/tday", "Low: trading days before elec", "Shape",
            "Retest vs pre-low %", "Retest drop from elec close %", "Retest days after elec",
            "Ret 1m %", "Ret 6m %", "6m worst low vs elec close %"]
    top = h.sort_values("Dist").head(5)[cols]
    now = pd.DataFrame([{"Metric": k, "Value": c[k]} for k in
                        ["Year", "Pre-low date", "Low: trading days before elec", "Rally low->elec close %",
                         "Rally slope %/tday", "N-height: peak vs low %", "Pullback peak->elec close %"]])
    pr = (h["Rally low->elec close %"] < c["Rally low->elec close %"]).mean() * 100
    now = pd.concat([now, pd.DataFrame([{"Metric": "% of past rallies smaller than current", "Value": round(pr, 1)}])])
    return now, top


def path_table(df):
    close, idx = df["close"], df.index
    offs = list(range(-PATH_PRE, PATH_POST + 1))
    cols = {}
    for y in MIDTERMS:
        e = election_date(y)
        if e > idx[-1] or idx[0] > e - pd.Timedelta(days=60):
            continue
        pe = idx.searchsorted(e, side="right") - 1
        base = close.iloc[pe]
        cols[y] = [close.iloc[pe + o] / base * 100 if 0 <= pe + o < len(close) else np.nan for o in offs]
    p = pd.DataFrame(cols, index=offs)
    out = pd.DataFrame({"Tday vs election": offs,
                        "Mean (elec=100)": p.mean(axis=1).round(2),
                        "Median (elec=100)": p.median(axis=1).round(2),
                        "% years above 100": (p.gt(100).sum(axis=1) / p.notna().sum(axis=1) * 100).round(1),
                        "N years": p.notna().sum(axis=1)})
    return out.join(p.round(2))


def autosize(ws):
    for col in ws.columns:
        L = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(L + 2, 10), 34)
    ws.column_dimensions["A"].width = 46


def main():
    results = {}
    with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
        summary_prob = []
        for name, fn in FILES.items():
            p = FOLDER / fn
            if not p.exists():
                print(f"!! missing {p}, skipping")
                continue
            df = load(p)
            print(f"{name}: {df.index[0].date()} -> {df.index[-1].date()}  ({len(df)} bars)")
            ev = pd.DataFrame([a for y in MIDTERMS if (a := analyze(df, y)) is not None])
            ev = add_shape(ev)
            ev.insert(0, "Index", name)
            results[name] = ev

            sheet = name
            row = 0
            ws_marks = []

            def put(tbl, title, index=False):
                nonlocal row
                pd.DataFrame([[title]]).to_excel(xw, sheet_name=sheet, startrow=row, index=False, header=False)
                ws_marks.append(row + 1)
                tbl.to_excel(xw, sheet_name=sheet, startrow=row + 1, index=index)
                start = row + 1
                row += len(tbl) + 4
                return start

            put(ev.drop(columns="Index"), f"{name} - EVENT TABLE (percent values are in %)")
            put(describe(ev), "AVERAGES / MEDIANS across completed midterms (descriptive only)")
            pr = prob_rows(ev)
            put(pr, "SIMPLE PROBABILITIES (completed midterms only; 2026 excluded)")
            put(cond_rows(ev), "CONDITIONAL PROBABILITIES (with tiny samples, treat as descriptive)")
            put(cycle_table(df), "4-YEAR CYCLE COMPARISON (same election-day anchor, every year in data)")
            now, top = analog_table(ev)
            if now is not None:
                put(now, "CURRENT (in-progress) MIDTERM STATUS")
                put(top, "5 CLOSEST HISTORICAL ANALOGS (by rally size and slope) - what happened next")
            pt = path_table(df)
            ps = put(pt, "AVERAGE PATH: index rebased to 100 at election-day close, by trading-day offset")
            ws = xw.sheets[sheet]
            for r in ws_marks:
                ws.cell(row=r, column=1).font = Font(bold=True, size=12, color="FFFFFF")
                for c in range(1, 12):
                    ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="305496")
            ch = LineChart()
            ch.title = f"{name}: avg midterm path (election = 100)"
            ch.height, ch.width = 9, 22
            data = Reference(ws, min_col=2 + 0, max_col=3, min_row=ps + 1, max_row=ps + 1 + len(pt))
            ch.add_data(Reference(ws, min_col=2, max_col=3, min_row=ps + 1, max_row=ps + 1 + len(pt)), titles_from_data=True)
            ch.set_categories(Reference(ws, min_col=1, min_row=ps + 2, max_row=ps + 1 + len(pt)))
            ws.add_chart(ch, f"{get_column_letter(len(pt.columns) + 3)}{ps}")
            autosize(ws)
            pr2 = pr.copy()
            pr2.insert(0, "Index", name)
            summary_prob.append(pr2)

        if summary_prob:
            allp = pd.concat(summary_prob)
            wide = allp.pivot_table(index="Observation", columns="Index", values="Probability %", sort=False)
            wide.reset_index().to_excel(xw, sheet_name="Summary", index=False)
            autosize(xw.sheets["Summary"])
            xw.book.move_sheet("Summary", offset=-len(xw.book.sheetnames) + 1)

    # ---------------- console output to paste back to me ----------------
    pd.set_option("display.width", 250, "display.max_columns", 40, "display.max_rows", 200)
    for name, ev in results.items():
        print("\n" + "=" * 30, name, "=" * 30)
        show = ["Year", "Pre-low date", "Low: trading days before elec", "Rally low->elec close %", "Rally slope %/tday",
                "Retest low date", "Retest vs pre-low %", "Retest drop from elec close %", "N width (tdays low->retest)",
                "Shape", "Retest low -> 6m %", "Ret 6m %", "6m worst low vs elec close %", "JanFeb low vs elec close %"]
        print(ev[show].to_string(index=False))
        print("\n", prob_rows(ev).to_string(index=False))
        print("\n", cond_rows(ev).to_string(index=False))
        now, top = analog_table(ev)
        if now is not None:
            print("\nCURRENT:\n", now.to_string(index=False))
            print("\nANALOGS:\n", top.to_string(index=False))
    print(f"\nSaved -> {OUT}")


if __name__ == "__main__":
    main()
